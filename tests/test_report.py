import io

import numpy as np
import pandas as pd

from report import Chart, Series, VLine, build_excel, build_pdf, label_rows


def test_excel_has_one_sheet_per_frame():
    xlsx = build_excel({"Özet": pd.DataFrame({"Kalem": ["Fiyat"], "Değer": ["₺10"]}),
                        "A very long sheet name that Excel would reject": pd.DataFrame({"x": [1.5, 2.5]})})
    book = pd.ExcelFile(io.BytesIO(xlsx))
    assert book.sheet_names == ["Özet", "A very long sheet name that Exc"]  # trimmed to 31 characters
    assert pd.read_excel(book, "Özet").iloc[0].tolist() == ["Fiyat", "₺10"]


def test_pdf_renders_turkish_text_and_lira_sign():
    x = np.linspace(10, 100, 500)
    chart = Chart("Kâr eğrisi", x, [Series("Net kâr", -(x - 50) ** 2 + 900, "#2a78d6")], xlabel="Fiyat (₺)",
                  vlines=[VLine(50, "Optimum ₺50", "#2a78d6"), VLine(30, "Mevcut ₺30", "#898781", dashed=True)],
                  band=(45, 55, "%90 aralık"))
    pdf = build_pdf("Fiyat Optimizasyon Motoru", "Şubat · ığüşöç İ", [("Optimum fiyat", "₺50,00")], chart,
                    "Notlar", ["Uyarı: talep eğrisi dışa doğru tahmin ediliyor."], "2026-10-26")
    assert pdf.startswith(b"%PDF") and len(pdf) > 5000


def test_excel_stores_formula_like_text_as_text():
    from openpyxl import load_workbook

    payload = '=HYPERLINK("http://example.com","click")'
    xlsx = build_excel({"Data": pd.DataFrame({"Product": [payload, "-5", "@SUM(A1)"], "Units": [1, 2, 3]})})
    ws = load_workbook(io.BytesIO(xlsx))["Data"]
    assert ws["A2"].data_type == "s" and ws["A2"].value == payload
    assert [ws[f"A{i}"].data_type for i in (3, 4)] == ["s", "s"]
    assert ws["B2"].value == 1


def test_pdf_with_long_summary_moves_chart_to_next_page():
    x = np.linspace(10, 100, 50)
    chart = Chart("Profit", x, [Series("Net profit", -(x - 50) ** 2 + 900, "#2a78d6")])
    pdf = build_pdf("Title", "Subtitle", [(f"Row {i}", "value") for i in range(40)], chart)
    assert pdf.count(b"/Type /Page") - pdf.count(b"/Type /Pages") == 2


def test_label_rows_stacks_only_overlapping_labels():
    # recommended, unconstrained and current labels starting close together, plus one far to the right
    spans = [(16.99, 19.5), (16.5, 19.0), (17.2, 19.4), (25.0, 27.0)]
    rows = label_rows(spans, gap=0.1)
    assert sorted(rows[:3]) == [0, 1, 2]  # the three close labels each get their own row
    assert rows[3] == 0  # the far one fits on the top row
    assert label_rows([]) == []


def test_pdf_with_overlapping_markers_builds():
    x = np.linspace(10, 30, 50)
    chart = Chart("Profit", x, [Series("Profit", -(x - 17) ** 2 + 100, "#2a78d6")],
                  vlines=[VLine(16.99, "Recommended $16.99", "#2a78d6"),
                          VLine(17.10, "Unconstrained $17.10", "#898781"),
                          VLine(16.89, "Current $16.89", "#898781", dashed=True),
                          VLine(29.9, "Edge $29.90", "#898781")])
    pdf = build_pdf("T", "", [("Price", "$16.99")], chart)
    assert pdf[:4] == b"%PDF"


def test_pdf_wraps_long_summary_values(monkeypatch):
    import report

    long = "KDV -9,83 ₺ · Ücretler -8,85 ₺ · Birim maliyet -23,72 ₺ · Marj 16,58 ₺ (%28,1)"
    heights = {}
    real = report._PDF.multi_cell

    def spy(self, w, h, text="", *args, **kwargs):
        y0 = self.get_y()
        out = real(self, w, h, text, *args, **kwargs)
        if not kwargs.get("dry_run") and text == long:
            heights[text] = self.get_y() - y0
        return out

    monkeypatch.setattr(report._PDF, "multi_cell", spy)
    x = np.linspace(10, 100, 50)
    chart = Chart("Kâr", x, [Series("Net kâr", -(x - 50) ** 2 + 900, "#2a78d6")])
    pdf = build_pdf("Başlık", "Alt", [("Birim başına", long), ("Sonraki", "x")], chart)
    assert pdf.startswith(b"%PDF")
    assert heights[long] == 12  # two 6 mm lines inside the value column, not one clipped line
